"""
Explore data/raw/gdx_data.xlsx and report what's actually in it.

Does not assume a fixed sheet layout: it lists every sheet, dumps the raw
first 10 rows as openpyxl sees them, then tries to auto-detect the header
row / date column / ticker columns so it can report a date range and row
count per ticker. If the auto-detection heuristics don't fit a given file,
the raw dump above still tells you what's there.
"""

import datetime
from pathlib import Path

import openpyxl
import pandas as pd

XLSX_PATH = Path("data/raw/gdx_data.xlsx")

MISSING_MARKERS = {None, "", "#N/A", "N/A", "NA", "n/a"}


def is_missing(value):
    return value in MISSING_MARKERS


def dump_raw_sheet(ws, n_rows=10):
    """Print the sheet's raw shape and its first n_rows exactly as stored."""
    print(f"  dimensions: {ws.dimensions}  (max_row={ws.max_row}, max_col={ws.max_column})")
    print(f"  first {n_rows} rows (1-indexed, raw cell values):")
    for i, row in enumerate(
        ws.iter_rows(min_row=1, max_row=min(n_rows, ws.max_row), values_only=True), start=1
    ):
        print(f"    row {i}: {row}")


def find_header_row(ws, max_scan=30):
    """First row with 2+ non-empty string cells -- our best guess at column labels."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(max_scan, ws.max_row), values_only=True), start=1):
        string_cells = [v for v in row if isinstance(v, str) and v.strip()]
        if len(string_cells) >= 2:
            return i, row
    return None, None


def find_date_column(ws, header_row_idx, sample_rows=50):
    """Among columns after the header row, pick the one whose values are mostly datetimes."""
    start = header_row_idx + 1
    end = min(start + sample_rows, ws.max_row)
    rows = list(ws.iter_rows(min_row=start, max_row=end, values_only=True))
    if not rows:
        return None
    n_cols = len(rows[0])
    best_col, best_count = None, 0
    for col in range(n_cols):
        count = sum(1 for r in rows if isinstance(r[col], (datetime.datetime, datetime.date)))
        if count > best_count:
            best_col, best_count = col, count
    return best_col if best_count > 0 else None


def analyze_sheet(ws, sheet_name):
    print(f"\n=== Sheet: {sheet_name!r} ===")
    dump_raw_sheet(ws, n_rows=10)

    header_row_idx, header_row = find_header_row(ws)
    if header_row_idx is None:
        print("  Could not auto-detect a header row (no row with 2+ text cells found).")
        return

    print(f"\n  Detected header row: row {header_row_idx} -> {header_row}")

    date_col = find_date_column(ws, header_row_idx)
    if date_col is None:
        print("  Could not auto-detect a date column near the header row.")
        return

    print(f"  Detected date column index (0-based): {date_col}")

    ticker_cols = {
        col_idx: name
        for col_idx, name in enumerate(header_row)
        if col_idx != date_col and isinstance(name, str) and name.strip()
    }
    if not ticker_cols:
        print("  No labeled ticker columns found alongside the date column.")
        return

    print(f"  Detected ticker columns: {ticker_cols}")

    # Data rows: everything after the header row, skipping any fully-blank rows.
    data_rows = [
        row
        for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=ws.max_row, values_only=True)
        if any(v is not None for v in row)
    ]

    print(f"\n  Date range / row count per ticker (non-missing values only):")
    for col_idx, name in ticker_cols.items():
        dates = []
        n_valid = 0
        n_missing = 0
        for row in data_rows:
            date_val = row[date_col] if date_col < len(row) else None
            price_val = row[col_idx] if col_idx < len(row) else None
            if not isinstance(date_val, (datetime.datetime, datetime.date)):
                continue
            if is_missing(price_val):
                n_missing += 1
                continue
            n_valid += 1
            dates.append(date_val)
        if dates:
            print(
                f"    {name}: {n_valid} rows with data, {n_missing} missing "
                f"({dates and min(dates)} to {dates and max(dates)})"
            )
        else:
            print(f"    {name}: no valid (non-missing) data rows found; {n_missing} missing/blank")


def main():
    if not XLSX_PATH.exists():
        raise SystemExit(f"File not found: {XLSX_PATH}")

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    print(f"File: {XLSX_PATH}")
    print(f"Sheet names: {wb.sheetnames}")

    for sheet_name in wb.sheetnames:
        analyze_sheet(wb[sheet_name], sheet_name)

    # Cross-check with pandas' own view of the raw grid (no header assumed).
    print("\n=== pandas cross-check (header=None, first 10 rows per sheet) ===")
    all_sheets = pd.read_excel(XLSX_PATH, sheet_name=None, header=None)
    for sheet_name, df in all_sheets.items():
        print(f"\n  Sheet {sheet_name!r}: shape={df.shape}")
        print(df.head(10).to_string())


if __name__ == "__main__":
    main()
